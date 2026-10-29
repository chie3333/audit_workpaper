"""재무제표·시산표 양식 엑셀: 양식 만들기, 업로드 읽기.

- 양식은 앱이 정한다(열 이름이 정확히 같아야 함). 회사 원본 양식의 열 매핑은 아직 없다.
- 금액은 원 단위 정수. 정상 잔액 방향(자산·비용은 차변, 부채·자본·수익은 대변)의 양수로 적고, 차감 계정은 음수로 적는다.
- 빈 금액은 빈칸 그대로 둔다. 0으로 바꾸지 않는다(계산은 engine이 막는다).
- 소계·합계 행은 넣지 않는다. 합계는 앱이 계산한다.
"""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import engine

STATEMENTS = ["재무상태표", "손익계산서"]
CATEGORIES = {"재무상태표": ["자산", "부채", "자본"], "손익계산서": ["수익", "비용"]}

FS_SHEET, TB_SHEET = "재무제표", "시산표"
FS_COLUMNS = ["재무제표", "분류", "표시 항목", "전기", "당기"]
TB_COLUMNS = ["계정코드", "계정과목", "표시 항목", "전기말", "당기"]

GUIDE = {
    "fs": [
        "회사가 작성한 재무제표(감사 전)를 표시 항목 단위로 한 행씩 적습니다.",
        "재무제표: 재무상태표 또는 손익계산서. 분류: 재무상태표는 자산·부채·자본, 손익계산서는 수익·비용.",
        "금액은 원 단위 정수, 양수로 적습니다. 소계·합계(자산총계, 영업이익, 당기순이익 등) 행은 넣지 않습니다.",
        "이익잉여금은 당기순이익을 포함한 재무상태표 금액을 적습니다.",
        "확인하지 않은 금액은 비워 둡니다. 빈칸은 0원으로 보지 않습니다.",
    ],
    "tb": [
        "회사 시산표(감사 전)를 계정과목 단위로 한 행씩 적습니다.",
        "표시 항목: 그 계정이 들어가는 재무제표 표시 항목 이름(재무제표 양식과 같은 이름). 비우면 앱에서 고를 수 있습니다.",
        "전기말: 재무상태표 계정은 전기말 잔액, 손익 계정은 전기 금액(비교용). 당기: 당기말 잔액(손익 계정은 당기 금액).",
        "금액은 정상 잔액 방향의 양수로 적습니다(자산·비용은 차변, 부채·자본·수익은 대변). 차감 계정(대손충당금 등)은 음수.",
        "이익잉여금 계정은 마감 전 잔액(당기순이익 대체 전)을 적습니다.",
        "감사인 수정분개는 이 파일에 넣지 않고 앱에서 입력합니다.",
    ],
}


class ImportError_(Exception):
    pass


# ---------------------------------------------------------------- 양식 만들기
def _sheet(wb: Workbook, title: str, columns: list[str], rows: list[list], guide: list[str], notice: str = ""):
    ws = wb.active
    ws.title = title
    head = PatternFill("solid", fgColor="DCE5EA")
    for i, name in enumerate(columns, 1):
        c = ws.cell(1, i, name)
        c.font, c.fill = Font(bold=True), head
        ws.column_dimensions[c.column_letter].width = 16 if i > 3 else 20
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            if isinstance(v, int):
                c.number_format = "#,##0"
    ws.freeze_panes = "A2"
    g = wb.create_sheet("작성 방법")
    g.column_dimensions["A"].width = 110
    for r, line in enumerate(([notice, ""] if notice else []) + guide, 1):
        g.cell(r, 1, line).alignment = Alignment(wrap_text=True)


def write_fs(rows: list[dict] | None = None, notice: str = "") -> bytes:
    wb = Workbook()
    data = [[r["statement"], r["category"], r["name"], r.get("prior"), r.get("current")] for r in rows or []]
    _sheet(wb, FS_SHEET, FS_COLUMNS, data, GUIDE["fs"], notice)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_tb(accounts: list[dict] | None = None, notice: str = "") -> bytes:
    wb = Workbook()
    data = [[a["code"], a["name"], a.get("fs_line") or None, a.get("prior"), a.get("unadjusted")] for a in accounts or []]
    _sheet(wb, TB_SHEET, TB_COLUMNS, data, GUIDE["tb"], notice)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- 업로드 읽기
def _rows(data: bytes, sheet: str, columns: list[str]):
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ImportError_("엑셀 파일(.xlsx)을 열 수 없음")
    if sheet not in wb.sheetnames:
        raise ImportError_(f"'{sheet}' 시트가 없음 — 앱에서 내려받은 양식을 쓰세요")
    it = wb[sheet].iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(it, [])][:len(columns)]
    if header != columns:
        raise ImportError_(f"첫 행의 열 이름이 양식과 다름: {header} (필요: {columns})")
    for n, row in enumerate(it, 2):
        row = list(row[:len(columns)]) + [None] * (len(columns) - len(row))
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        yield n, row


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _amount(v, n, label, errors) -> int | str:
    """정수 금액. 빈칸은 ''(미입력)로 둔다."""
    val, err = engine.parse_amount(v)
    if err:
        errors.append(f"{n}행 {label}: {err} ('{v}')")
        return ""
    return "" if val is None else val


def parse_fs(data: bytes) -> dict:
    rows, errors, seen = [], [], set()
    for n, (stmt, cat, name, prior, cur) in _rows(data, FS_SHEET, FS_COLUMNS):
        stmt, cat, name = _text(stmt), _text(cat), _text(name)
        if stmt not in STATEMENTS:
            errors.append(f"{n}행: 재무제표는 {'·'.join(STATEMENTS)} 중 하나 ('{stmt}')")
            continue
        if cat not in CATEGORIES[stmt]:
            errors.append(f"{n}행: {stmt}의 분류는 {'·'.join(CATEGORIES[stmt])} 중 하나 ('{cat}')")
            continue
        if not name:
            errors.append(f"{n}행: 표시 항목 이름이 없음")
            continue
        if (stmt, name) in seen:
            errors.append(f"{n}행: 같은 표시 항목이 두 번 있음 ('{name}')")
            continue
        seen.add((stmt, name))
        rows.append({"statement": stmt, "category": cat, "name": name, "loc": f"{FS_SHEET}!{n}행",
                     "prior": _amount(prior, n, "전기", errors), "current": _amount(cur, n, "당기", errors)})
    if not rows and not errors:
        errors.append("표시 항목 행이 없음")
    return {"rows": rows, "errors": errors}


def parse_tb(data: bytes) -> dict:
    accts, errors, seen = [], [], set()
    for n, (code, name, line, prior, cur) in _rows(data, TB_SHEET, TB_COLUMNS):
        code, name = _text(code), _text(name)
        if isinstance(code, str) and code.endswith(".0"):
            code = code[:-2]
        if not code or not name:
            errors.append(f"{n}행: 계정코드와 계정과목이 모두 있어야 함")
            continue
        if code in seen:
            errors.append(f"{n}행: 계정코드가 두 번 있음 ('{code}')")
            continue
        seen.add(code)
        accts.append({"code": code, "name": name, "fs_line": _text(line), "loc": f"{TB_SHEET}!{n}행",
                      "prior": _amount(prior, n, "전기말", errors), "unadjusted": _amount(cur, n, "당기", errors)})
    if not accts and not errors:
        errors.append("계정과목 행이 없음")
    return {"accounts": accts, "errors": errors}
