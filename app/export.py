"""엑셀 조서 생성.

- 서식 엑셀의 셀을 채우지 않고 새 양식으로 만든다.
- 기본은 한 시트. 검토표 행이 기준을 넘으면 상세 시트로 분리하고 본 시트에서 수식으로 연결한다.
- 차이·합계는 엑셀 수식으로 쓴다. 값은 engine.compute와 같아야 하며 테스트에서 대조한다.
- 증거 파일은 포함하지 않는다. 자료 식별정보와 위치만 기록한다.
"""
from __future__ import annotations

import io
import math
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import engine

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FT = Font(name=FONT, size=15, bold=True)
FS = Font(name=FONT, size=9, color="5A6670")
FW = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FR = Font(name=FONT, size=10, bold=True, color="9C2B1E")
SEC = PatternFill("solid", fgColor="23404F")
HEAD = PatternFill("solid", fgColor="DCE5EA")
KEY = PatternFill("solid", fgColor="F1F4F6")
WARN = PatternFill("solid", fgColor="FBE3C7")
SAMPLE = PatternFill("solid", fgColor="F6D7D2")
OK = PatternFill("solid", fgColor="DDEEDD")
BLANKFILL = PatternFill("solid", fgColor="FFF4C2")
thin = Side(style="thin", color="B9C4CB")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
NUM = "#,##0;-#,##0;0"
WIDTHS = {"A": 5, "B": 14, "C": 18, "D": 13, "E": 16, "F": 16, "G": 17, "H": 22, "I": 22, "J": 26}
CHARS_PER_LINE = 70  # C:J 병합 영역에 한글 기준 대략적인 한 줄 글자 수

DOC_LABEL = {"schedule": "명세서", "subledger": "보조부", "gl": "총계정원장",
             "terms": "증서·약정서", "restriction": "사용제한 자료", "other": "기타"}
DATE = "yyyy-mm-dd"
CLASSIFY = "classification_review"


def _amount_cell(raw):
    v, err = engine.parse_amount(raw)
    if err:
        return str(raw)
    return v


def _ref(content: dict, ref: str, detail: bool = False) -> str:
    """출처 참조(예: 'KSA2026:330.20')를 사람이 읽는 이름으로. 출처 이름은 콘텐츠의 출처 목록에서 가져온다."""
    sid, loc = ref.split(":", 1)
    src = content["sources"].get(sid, {})
    head = "감사기준서" if sid == "KSA2026" else "서식"
    label = src.get("locations", {}).get(loc, "") if detail else ""
    return f"{head} {loc}" + (f" ({label})" if label else "")


def _form_text(content: dict) -> str:
    """이 절차의 서식 위치. 계정 설정의 form_refs가 없으면(이전 콘텐츠 버전) form_ref 하나를 쓴다."""
    if "form_refs" in content:
        refs = content["form_refs"]
        if not refs:
            return "미확인 — " + (content.get("form_note") or "서식에서 이 절차의 위치를 확인하지 못함")
        title = content["sources"].get(refs[0].split(":")[0], {}).get("title", "")
        return f"{title} — " + " · ".join(_ref(content, r, detail=True) for r in refs)
    return _ref(content, content["form_ref"], detail=True) if content.get("form_ref") else "미확인"


def _assertion_source(content: dict) -> str:
    a = content["assertions"]
    if a.get("refs"):
        return "[출처: " + ", ".join(_ref(content, r) for r in a["refs"]) + "]"
    decided = (a.get("review") or {}).get("status") == "사용자 검토 완료"
    return "[출처 미확인" + (" — 사용자가 정함" if decided else "") + "]"


def _basis_text(content: dict) -> str:
    """근거: 계정 설정의 basis(근거매핑에서 확인한 위치만)와 미확인 메모."""
    b = content.get("basis")
    if not b:
        return _form_text(content) + "  (근거 목록 없음 — 이전 콘텐츠 버전)"
    refs = b.get("refs", [])
    text = " · ".join(_ref(content, r) for r in refs)
    ksa = content["sources"].get("KSA2026", {})
    if any(r.startswith("KSA2026:") for r in refs):
        text += f" ({ksa.get('title', '')} {ksa.get('edition', '')})"
    if b.get("note"):
        text += (". " if text else "") + b["note"]
    return text or "미확인"


def _won(v) -> str:
    return "미입력" if v is None else f"{v:,}"


def _chg(amount, rate, prior) -> str:
    """증감액·증감률. 증가 ▲, 감소 ▼ (계산은 engine)."""
    if amount is None:
        return "증감 미입력"
    mark = "▲ " if amount > 0 else "▼ " if amount < 0 else ""
    r = f"{abs(rate):.1f}%" if rate is not None else ("신규" if prior == 0 and amount else "—")
    rmark = mark if rate is not None else ""
    return f"증감 {mark}{abs(amount):,} ({rmark}{r})"


def _lines(text: str, per_line: int = CHARS_PER_LINE) -> int:
    if not text:
        return 1
    return sum(max(1, math.ceil(len(part) / per_line)) for part in str(text).split("\n"))


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.r = 1
        for col, w in WIDTHS.items():
            ws.column_dimensions[col].width = w

    def section(self, title):
        self.r += 1
        ws = self.ws
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=10)
        c = ws.cell(self.r, 1, title)
        c.font, c.fill = FW, SEC
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[self.r].height = 20
        self.r += 1

    def kv(self, key, value, fill=None, font=None):
        ws = self.ws
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=2)
        k = ws.cell(self.r, 1, key)
        k.font, k.fill, k.alignment = FB, KEY, WRAP
        ws.merge_cells(start_row=self.r, start_column=3, end_row=self.r, end_column=10)
        text = "" if value is None else str(value)
        v = ws.cell(self.r, 3, text)
        v.font, v.alignment = (font or F), WRAP
        if fill:
            v.fill = fill
        for col in range(1, 11):
            ws.cell(self.r, col).border = BOX
        ws.row_dimensions[self.r].height = 15 * _lines(text)
        self.r += 1

    def note(self, text, font=FS, fill=None):
        ws = self.ws
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=10)
        c = ws.cell(self.r, 1, text)
        c.font, c.alignment = font, WRAP
        if fill:
            c.fill = fill
        ws.row_dimensions[self.r].height = 15 * _lines(text, 110)
        self.r += 1

    def header(self, labels, start_col=1):
        for i, lab in enumerate(labels):
            c = self.ws.cell(self.r, start_col + i, lab)
            c.font, c.fill, c.border = FB, HEAD, BOX
            c.alignment = Alignment(wrap_text=True, vertical="center")
        self.ws.row_dimensions[self.r].height = 28
        self.r += 1

    def row(self, values, start_col=1, height_cols=None):
        mx = 1
        for i, v in enumerate(values):
            c = self.ws.cell(self.r, start_col + i, v)
            c.font, c.border, c.alignment = F, BOX, WRAP
            if isinstance(v, str) and height_cols and (start_col + i) in height_cols:
                mx = max(mx, _lines(v, max(8, int(WIDTHS[get_column_letter(start_col + i)] * 0.55))))
        self.ws.row_dimensions[self.r].height = 15 * mx
        self.r += 1
        return self.r - 1


def _recon_table(sh: Sheet, st: dict, content: dict) -> dict:
    """대사 행을 쓰고 행 범위를 돌려준다. 수식 셀 위치는 요약표에서 참조한다."""
    rlab = content["row"]
    extra = rlab.get("extra_columns", [])
    sh.header(["#", rlab["id_label"], rlab["name_label"], extra[0]["label"] if extra else "구분",
               "명세서 금액", "보조부 금액", "차이\n(명세서−보조부)", "명세서 근거 위치", "보조부 근거 위치", "비고·비경상 표시"])
    first = sh.r
    for i, row in enumerate(st.get("rows", []), 1):
        r = sh.r
        note = row.get("note") or ""
        if row.get("unusual"):
            note = (note + "\n" if note else "") + f"[비경상] {row.get('unusual_reason') or '(이유 미기록)'}"
        sh.row([i, row.get("item_id") or "", row.get("name") or "", row.get(extra[0]["key"], "") if extra else "",
                _amount_cell(row.get("schedule_amount")), _amount_cell(row.get("subledger_amount")),
                f'=IF(OR(E{r}="",F{r}=""),"미입력",IF(OR(ISTEXT(E{r}),ISTEXT(F{r})),"형식 오류",E{r}-F{r}))',
                row.get("schedule_loc") or "", row.get("subledger_loc") or "", note], height_cols={8, 9, 10})
        for col in (5, 6, 7):
            sh.ws.cell(r, col).number_format = NUM
            sh.ws.cell(r, col).alignment = Alignment(horizontal="right", vertical="top")
    last = sh.r - 1
    if last >= first:
        rng = f"E{first}:F{last}"
        sh.ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISBLANK(E{first})'], fill=BLANKFILL))
    return {"first": first, "last": last, "n": last - first + 1 if last >= first else 0}


def _summary(sh: Sheet, rng: dict, st: dict, gl_raw, gl_loc: str, ref_sheet: str | None) -> dict:
    """요약 수식. ref_sheet가 있으면 상세 시트의 행을 참조한다."""
    pre = f"'{ref_sheet}'!" if ref_sheet else ""
    if rng["n"]:
        e = f"{pre}E{rng['first']}:E{rng['last']}"
        f = f"{pre}F{rng['first']}:F{rng['last']}"
        recalc = f'=IF(COUNT({e})<ROWS({e}),"미입력",SUM({e}))'
        subt = f'=IF(COUNT({f})<ROWS({f}),"미입력",SUM({f}))'
    else:
        recalc = subt = "미입력"
    ws = sh.ws
    cells = {}

    def line(label, col, value, loc="", fmt=True, bold=False):
        r = sh.r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        k = ws.cell(r, 1, label)
        k.font, k.fill, k.border = FB if bold else F, KEY, BOX
        for c in range(1, 11):
            ws.cell(r, c).border = BOX
        c = ws.cell(r, col, value)
        c.font = FB if bold else F
        c.alignment = Alignment(horizontal="right")
        if fmt:
            c.number_format = NUM
        if loc:
            ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=10)
            lc = ws.cell(r, 8, loc)
            lc.font, lc.alignment = FS, WRAP
        sh.r += 1
        return f"{get_column_letter(col)}{r}"

    cells["recalc"] = line("명세서 재계산 합계", 5, recalc)
    stated = _amount_cell(st.get("stated_total"))
    cells["stated"] = line("명세서 표시 합계", 5, stated, st.get("stated_total_loc") or "")
    if stated is None:
        ws[cells["stated"]].fill = BLANKFILL
    s_, r_ = cells["stated"], cells["recalc"]
    cells["stated_vs_recalc"] = line("표시 합계 − 재계산 합계", 7,
                                     f'=IF(OR({s_}="",ISTEXT({s_}),ISTEXT({r_})),"미입력",{s_}-{r_})', bold=True)
    cells["sub_total"] = line("보조부 합계", 6, subt)
    gl = _amount_cell(gl_raw)
    cells["gl"] = line("총계정원장 잔액", 6, gl, gl_loc or "")
    if gl is None:
        ws[cells["gl"]].fill = BLANKFILL
    g_, b_ = cells["gl"], cells["sub_total"]
    cells["subledger_vs_gl"] = line("보조부 합계 − 총계정원장 잔액", 7,
                                    f'=IF(OR({g_}="",ISTEXT({g_}),ISTEXT({b_})),"미입력",{b_}-{g_})', bold=True)
    return cells


def _date_cell(raw):
    d, err = engine.parse_date(raw)
    if err:
        return str(raw)
    return d


def _classify_tables(sh: Sheet, data: dict, content: dict, fields: dict) -> dict:
    """분류 검토표 두 개(기간·분류)와 분류별 합계. 일수·재분류 금액·합계는 수식으로 쓴다."""
    ws = sh.ws
    rlab = content["row"]
    extra = rlab.get("extra_columns", [])
    classes = content["classify"]["classes"]
    rows = data["classify"].get("rows", [])
    cellmap = {"rows": {}, "totals": {}}

    r = sh.r
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    k = ws.cell(r, 1, "대상일(기준일) — 대상일→만기 일수의 기준")
    k.font, k.fill = FB, KEY
    asof = ws.cell(r, 5, _date_cell(fields.get("as_of_date")))
    asof.number_format, asof.font = DATE, FB
    for c in range(1, 11):
        ws.cell(r, c).border = BOX
    asof_ref = f"$E${r}"
    sh.r += 1

    sh.note("(1) 만기와 기간 — 일수는 날짜로만 계산했다. 몇 개월 이내인지, 어느 분류가 맞는지는 프로그램이 판단하지 않는다.", font=FB)
    sh.header(["#", rlab["id_label"], rlab["name_label"], extra[0]["label"] if extra else "구분",
               "금액", "취득일", "만기일", "취득→만기\n(일)", "대상일→만기\n(일)", "날짜 근거 위치"])
    for i, row in enumerate(rows, 1):
        r = sh.r
        no_mat = row.get("maturity_type") == "만기 없음"
        mat = "만기 없음" if no_mat else _date_cell(row.get("maturity_date"))
        acq = "" if no_mat else _date_cell(row.get("acquired_date"))
        sh.row([i, row.get("item_id") or "", row.get("name") or "", row.get(extra[0]["key"], "") if extra else "",
                _amount_cell(row.get("amount")), acq, mat,
                f'=IF(G{r}="만기 없음","해당 없음",IF(OR(F{r}="",G{r}=""),"미입력",IF(OR(ISTEXT(F{r}),ISTEXT(G{r})),"형식 오류",G{r}-F{r})))',
                f'=IF(G{r}="만기 없음","해당 없음",IF(OR(G{r}="",{asof_ref}=""),"미입력",IF(OR(ISTEXT(G{r}),ISTEXT({asof_ref})),"형식 오류",G{r}-{asof_ref})))',
                row.get("terms_loc") or ""], height_cols={10})
        ws.cell(r, 5).number_format = NUM
        for col in (6, 7):
            ws.cell(r, col).number_format = DATE
        for col in (5, 8, 9):
            ws.cell(r, col).alignment = Alignment(horizontal="right", vertical="top")
            if col != 5:
                ws.cell(r, col).number_format = NUM
        if ws.cell(r, 5).value is None:
            ws.cell(r, 5).fill = BLANKFILL
        cellmap["rows"][row["id"]] = {"acq_days": f"H{r}", "remain_days": f"I{r}"}
    sh.r += 1

    sh.note("(2) 사용제한과 분류 — 감사인 판단 분류와 근거는 수행자가 직접 적은 내용이다.", font=FB)
    sh.header(["#", rlab["id_label"], "사용제한", "제한 내용·근거 위치", "금액", "회사 분류", "감사인 판단 분류",
               "재분류 금액", "판단 근거", "금액 근거 위치"])
    first = sh.r
    for i, row in enumerate(rows, 1):
        r = sh.r
        rs = row.get("restricted") or "미확인"
        rnote = "\n".join(x for x in (row.get("restriction_note"), row.get("restriction_loc")) if x)
        sh.row([i, row.get("item_id") or "", rs, rnote, _amount_cell(row.get("amount")),
                row.get("company_class") or None, row.get("auditor_class") or None,
                f'=IF(OR(E{r}="",F{r}="",G{r}=""),"미입력",IF(ISTEXT(E{r}),"형식 오류",IF(F{r}=G{r},0,E{r})))',
                row.get("basis") or "", row.get("amount_loc") or ""], height_cols={4, 9, 10})
        for col in (5, 8):
            ws.cell(r, col).number_format = NUM
            ws.cell(r, col).alignment = Alignment(horizontal="right", vertical="top")
        if rs == "미확인":
            ws.cell(r, 3).fill = WARN
        for col in (5, 6, 7):
            if ws.cell(r, col).value is None:
                ws.cell(r, col).fill = BLANKFILL
        cellmap["rows"][row["id"]]["reclass"] = f"H{r}"
    last = sh.r - 1
    sh.r += 1

    sh.note("(3) 분류별 합계 — 금액이나 분류가 빈 행이 있으면 합계를 '미입력'으로 표시한다.", font=FB)
    hr = sh.r
    sh.header(["", "분류", "", "", "회사 분류 합계", "감사인 판단 합계", "차이\n(판단−회사)", "", "", ""])
    ws.merge_cells(start_row=hr, start_column=2, end_row=hr, end_column=4)
    n = last - first + 1
    for k in classes:
        r = sh.r
        if n > 0:
            amt, cc, ac = f"E{first}:E{last}", f"F{first}:F{last}", f"G{first}:G{last}"
            bad = f"COUNTBLANK({amt})+SUMPRODUCT(--ISTEXT({amt}))"
            comp = f'=IF({bad}+COUNTBLANK({cc})>0,"미입력",SUMIF({cc},"{k}",{amt}))'
            aud = f'=IF({bad}+COUNTBLANK({ac})>0,"미입력",SUMIF({ac},"{k}",{amt}))'
        else:
            comp = aud = "미입력"
        sh.row(["", k, "", "", comp, aud, f'=IF(OR(ISTEXT(E{r}),ISTEXT(F{r})),"미입력",F{r}-E{r})', "", "", ""])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        for col in (5, 6, 7):
            ws.cell(r, col).number_format = NUM
            ws.cell(r, col).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(r, 7).font = FB
        cellmap["totals"][k] = {"company": f"E{r}", "auditor": f"F{r}", "diff": f"G{r}"}
    return cellmap


def build(wp: dict, computed: dict | None = None) -> tuple[bytes, dict]:
    """엑셀 바이트와, 테스트용 셀 위치 지도를 돌려준다."""
    content, data = wp["content"], wp["data"]
    computed = computed or engine.compute(content, data, wp.get("ledger"))
    fields = data["fields"]
    mats = {m["id"]: m for m in data.get("materials", [])}
    names = {a["code"]: a["name"] for a in content["assertion_list"]}
    threshold = content.get("export", {}).get("detail_sheet_row_threshold", 30)
    status = computed["status"]
    open_items = [i for i in computed["open_items"] if i["level"] in ("block", "warn")]

    wb = Workbook()
    ws = wb.active
    ws.title = "조서"
    ws.sheet_view.showGridLines = False
    sh = Sheet(ws)
    cellmap = {"sets": {}}
    is_cls = content["template"]["id"] == CLASSIFY

    ws.merge_cells("A1:J1")
    ws["A1"] = content["title"]
    ws["A1"].font = FT
    ws.row_dimensions[1].height = 24
    ws.merge_cells("A2:J2")
    ws["A2"] = f"{wp['company']} · 대상일 {fields.get('as_of_date') or wp['period_end']} · 조서번호 {fields.get('wp_no') or '(미입력)'}"
    ws["A2"].font = FS
    sh.r = 3
    if wp.get("is_sample"):
        sh.note("가상 데이터 — 회사·자료·금액·수행 사실이 모두 가상이며 실제 감사 수행 기록이 아님", font=FR, fill=SAMPLE)
    flag = []
    if status["review"] != "검토 완료":
        flag.append(f"검토: {status['review']}")
    if status["findings_unresolved"]:
        flag.append(f"미처리 발견 사항 {status['findings_unresolved']}건")
    blocks = [i for i in open_items if i["level"] == "block"]
    if blocks:
        flag.append(f"남은 작업 {len(blocks)}건")
    sh.note(f"수행 상태: {status['procedure']}  |  발견 사항: {status['findings']}  |  검토 상태: {status['review']}",
            font=FB, fill=WARN if flag else OK)
    if flag:
        sh.note("주의: " + " · ".join(flag) + " — 아래 7번 '남은 작업'에 목록이 있음", font=FR)

    # 1. 기본정보
    sh.section("1. 기본정보")
    sh.kv("회사", wp["company"])
    sh.kv("보고기간 종료일", wp["period_end"])
    sh.kv("통화·단위", f"{wp['currency']} · {wp['unit']}")
    sh.kv("계정", content["account"])
    sh.kv("절차", content["title"])
    sh.kv("서식 위치", _form_text(content))
    sh.kv("콘텐츠 버전", f"{content['content_key']} (해시 {content['hash'][:12]})")
    sh.kv("조서 저장본", f"리비전 {wp['rev']} · {wp['saved_at']} · {wp['saved_by']}" + ("  ※ 최신 저장본 아님" if wp["rev"] != wp["latest_rev"] else ""))
    sh.kv("편집 잠금", "잠김 (업무용 잠금이며 최종 감사파일 취합·법정 보존 기능이 아님)" if wp["locked"] else "잠기지 않음")
    sh.kv("출력 일시", datetime.now().strftime("%Y-%m-%d %H:%M"))

    # 2. 목적과 범위
    sh.section("2. 목적과 범위")
    sh.kv("목적", fields.get("purpose"))
    sh.kv("검사 범위", fields.get("scope"))
    applied = data["assertions"].get("applied", [])
    default = content["assertions"]["default"]
    a_text = ", ".join(f"{names.get(c, c)}({c})" for c in applied) or "(없음)"
    a_text += "  — 콘텐츠 기본값과 같음" if sorted(applied) == sorted(default) else \
        f"  — 콘텐츠 기본값({', '.join(default)})에서 변경. 변경 이유: {data['assertions'].get('change_reason') or '(미기록)'}"
    a_text += "  " + _assertion_source(content)
    sh.kv("적용 경영진 주장", a_text)
    sh.kv("관련 위험·계획 조서", fields.get("risk_ref"))
    sh.kv("중요성 참조", fields.get("materiality_ref"))
    sh.kv("관련 조서", fields.get("related_wps"))
    link = computed.get("linked")
    if link is not None:
        if not link["available"]:
            sh.kv("정산표 계정과목", "이 감사건에 정산표 없음")
        elif not link["rows"]:
            sh.kv("정산표 계정과목", "(선택하지 않음)", fill=WARN)
        else:
            txt = "\n".join(f"{r['code']} {r['name']} ({r['fs_line']}) — 전기말 {_won(r['prior'])} · 수정후 {_won(r['adjusted'])} · "
                            f"{_chg(r['change'], r['change_rate'], r['prior'])}" for r in link["rows"])
            txt += (f"\n합계 — 전기말 {_won(link['prior'])} · 수정후 {_won(link['adjusted'])} · "
                    f"{_chg(link['change'], link['change_rate'], link['prior'])}  (참조값. 조서 입력값을 대신하지 않음)")
            sh.kv("정산표 계정과목", txt)
    sh.kv("근거", _basis_text(content))

    # 3. 사용 자료
    sh.section("3. 사용 자료")
    sh.note("증거 파일은 이 엑셀에 포함되어 있지 않다. 자료의 식별정보와 위치만 기록했다. '등록만 됨'은 자료를 받았지만 증거로 확인한 기록이 없다는 뜻이다.")
    sh.header(["", "구분", "자료명", "버전·원본/수정본", "제공자", "입수 상태·입수일", "기준일", "위치(파일·시트 등)", "증거 확인", "남은 작업·비고"])
    for m in data.get("materials", []):
        kind = m.get("kind") or "원본"
        if kind == "수정본":
            kind += f" (원본: {mats.get(m.get('original_id'), {}).get('name', '미연결')})"
        rec = m.get("receipt") or "상태 미입력"
        if m.get("received_date"):
            rec += f" {m['received_date']}"
        if m.get("confirmed"):
            conf = f"확인함\n{m.get('confirmed_by') or '?'} · {m.get('confirmed_date') or '?'}\n범위: {m.get('confirmed_scope') or '(미기록)'}"
        else:
            conf = "등록만 됨" if m.get("receipt") == "입수" else "—"
        r = sh.row(["", DOC_LABEL.get(m.get("role")) or content["documents"].get(m.get("role"), m.get("role")), m.get("name") or "",
                    f"{m.get('version_label') or ''}\n{kind}", m.get("provider") or "", rec, m.get("as_of_date") or "",
                    m.get("location") or "", conf, m.get("remaining_work") or m.get("note") or ""], height_cols={3, 4, 8, 9, 10})
        if m.get("receipt") != "입수":
            ws.cell(r, 6).fill = WARN
        if m.get("receipt") == "입수" and not m.get("confirmed"):
            ws.cell(r, 9).fill = WARN

    # 4. 수행 내역
    sh.section("4. 수행 내역")
    ps = fields.get("procedure_status") or "미착수"
    sh.kv("절차 수행 상태", ps)
    if ps == "해당 없음":
        sh.kv("해당 없음 사유", fields.get("na_reason"))
    if ps == "대체절차 수행":
        sh.kv("수행한 대체절차", fields.get("alt_procedure"))
        sh.kv("목적 달성 방법", fields.get("alt_how"))
        sh.kv("이탈 이유", fields.get("deviation_reason"))
    sh.kv("실제 수행 내역", fields.get("work_performed"))
    if is_cls:
        sh.section("5. 분류 검토표")
        sh.note("빈칸은 확인하지 않은 값이다(노란 칸). 금액이나 분류가 빈 행이 있으면 분류별 합계를 계산하지 않는다.")
        cellmap["classify"] = _classify_tables(sh, data, content, fields)
    else:
        lc = fields.get("list_check") or "미수행"
        sh.kv("목록 양방향 대조", lc + (f"\n{fields.get('list_check_note')}" if fields.get("list_check_note") else ""))
        sc = data.get("scan", {})
        sct = sc.get("status") or "미결정"
        if sct == "수행":
            sct += f"\n{sc.get('note') or ''}"
        elif sct == "해당 없음":
            sct += f" — 사유: {sc.get('na_reason') or '(미기록)'}"
        sh.kv("명세서 스캔(선택 단계)", sct)

        # 5. 검토표
        sh.section("5. 검토표 · 계산검증과 대사")
        sh.note("빈칸은 확인하지 않은 금액이며 0원으로 계산하지 않는다(노란 칸). 빈칸이 있으면 합계와 차이는 '미입력'으로 표시된다.")
        sets = data["recon"]["sets"]
        detail = None
        for idx, st in enumerate(reversed(sets)):
            is_active = idx == 0
            mat = mats.get(st.get("schedule_material_id"), {})
            title = f"{'현재 대사' if is_active else '이전 대사 기록(수정 전, 보존용)'}: {st.get('label') or ''} — 명세서 자료: {mat.get('name') or '(미연결)'} {mat.get('version_label') or ''}"
            sh.note(title, font=FB)
            n = len(st.get("rows", []))
            if n > threshold:
                if detail is None:
                    detail = Sheet(wb.create_sheet("검토표_상세"))
                    detail.ws.sheet_view.showGridLines = False
                detail.note(f"{wp['company']} · {content['title']} — {title}", font=FB)
                rng = _recon_table(detail, st, content)
                detail.r += 1
                sh.note(f"항목 {n}건은 '검토표_상세' 시트에 있다. 아래 합계는 상세 시트를 수식으로 참조한다.")
                cells = _summary(sh, rng, st, data["recon"].get("gl_balance"), data["recon"].get("gl_loc"), "검토표_상세")
                cellmap["sets"][st["id"]] = {"sheet_rows": "검토표_상세", "rows": rng, "summary": cells}
            else:
                rng = _recon_table(sh, st, content)
                cells = _summary(sh, rng, st, data["recon"].get("gl_balance"), data["recon"].get("gl_loc"), None)
                cellmap["sets"][st["id"]] = {"sheet_rows": "조서", "rows": rng, "summary": cells}
            sh.r += 1

    # 6. 발견 사항
    sh.section("6. 발견 사항과 후속 대응")
    findings = data.get("findings", [])
    if not findings:
        sh.note("기록된 발견 사항 없음")
    else:
        for i, f in enumerate(findings, 1):
            et = f.get("explanation_type") or "없음"
            if et == "담당자 설명만":
                basis = f"담당자 설명만 받음 — {f.get('person') or '?'} · {f.get('discussed_at') or '?'}\n{f.get('explanation') or ''}"
            elif et == "자료로 재확인":
                basis = f"자료로 재확인: {mats.get(f.get('recheck_material_id'), {}).get('name', '(미연결)')}"
            else:
                basis = "설명·재확인 기록 없음"
            done = f.get("status") == "처리 완료"
            sh.note(f"발견 사항 {i}. {f.get('title') or '(제목 없음)'}  [{f.get('status') or '상태 미입력'}]", font=FB, fill=None if done else WARN)
            amt = f.get("amount")
            sh.kv("금액", f"{amt:,}원" if isinstance(amt, int) else (amt or ""))
            sh.kv("사실", f.get("fact"))
            sh.kv("설명·재확인 근거", basis)
            sh.kv("원인", f.get("cause"))
            sh.kv("후속 확인", f.get("followup"))
            if done:
                sh.kv("처리 내용", f.get("resolution"))
                sh.kv("처리 근거", f.get("resolution_basis"))
            sh.r += 1

    # 7. 결론
    sh.section("7. 결론과 미해결 사항")
    if is_cls:
        sh.kv("이번 조서의 분류 검토 결과", fields.get("classify_result"))
    else:
        sh.kv("이번 조서의 대사 결과", fields.get("recon_result"))
        fl = fields.get("fs_link_status") or "(미입력)"
        sh.kv("재무제표 정보와 회계기록의 연결", fl + (f" — 참조: {fields.get('fs_link_ref')}" if fl == "별도 조서 참조" else ""),
              fill=WARN if fl in ("확인 중", "미수행", "(미입력)") else None)
        if fl in ("일치 확인", "조정 후 확인", "별도 조서 참조"):
            sh.kv("연결 근거", fields.get("fs_link_basis"))
            sh.kv("명시 문장", fields.get("fs_link_statement"))
        else:
            sh.kv("명시 문장", f"작성하지 않음 — 연결 검토 상태: {fl}")
    sh.kv("결론", fields.get("conclusion"))
    if open_items:
        sh.note(f"출력 시점 남은 작업 {len(blocks)}건, 주의 {len(open_items) - len(blocks)}건", font=FR)
        hr = sh.r
        sh.header(["", "구분", "단계", "내용", "", "", "", "", "", ""])
        ws.merge_cells(start_row=hr, start_column=4, end_row=hr, end_column=10)
        step_title = {s["id"]: s["title"] for s in content["steps"]}
        for it in open_items:
            r = sh.r
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
            sh.row(["!" if it["level"] == "block" else "주의", it["category"], step_title.get(it["step"], it["step"]), it["message"]])
            ws.row_dimensions[r].height = 15 * _lines(it["message"], 95)
            if it["level"] == "block":
                ws.cell(r, 1).fill = WARN
    else:
        sh.note("출력 시점 남은 작업 없음")

    # 8. 작성·검토 기록
    sh.section("8. 작성·검토 기록")
    sh.kv("수행자", fields.get("performer"))
    sh.kv("수행 종료일", fields.get("performed_date"))
    rs = fields.get("review_status") or "미검토"
    sh.kv("검토 상태", rs, fill=None if rs == "검토 완료" else WARN, font=None if rs == "검토 완료" else FR)
    if rs != "미검토":
        sh.kv("검토자", fields.get("reviewer"))
        sh.kv("검토일", fields.get("review_date"))
        sh.kv("검토 범위", fields.get("review_scope"))
    comments = data.get("review_comments", [])
    if comments:
        sh.header(["#", "작성자", "일자", "검토 의견", "", "", "답변", "", "처리", ""])
        for i, c in enumerate(comments, 1):
            r = sh.r
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
            ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
            sh.row([i, c.get("by") or "", c.get("date") or "", c.get("text") or "", "", "", c.get("response") or "", "",
                    "처리됨" if c.get("resolved") else "미처리", ""])
    sh.note("엑셀 다운로드는 검토 완료를 뜻하지 않는다. 검토 상태는 위 기록을 따른다.")
    sh.note(f"콘텐츠 {content['content_key']} · 해시 {content['hash'][:12]} · 리비전 {wp['rev']}")

    ws.freeze_panes = "A3"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), cellmap
