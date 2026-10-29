"""재무제표·시산표 업로드 → 정산표(수정분개 포함) → 계정원장 → 분개장, 그리고 조서의 계정과목 참조."""
import copy
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import content, engine, export, imports, records, samples
from app.main import app


def _ledger():
    return copy.deepcopy(samples.ledger())


def _acct(L, code):
    return next(a for a in L["accounts"] if a["code"] == code)


def _msgs(issues, level=None):
    return [i["message"] for i in issues if level is None or i["level"] == level]


# ---------------------------------------------------------------- 정산표 계산
def test_sample_ledger_ties_out():
    c = engine.compute_ledger(_ledger())
    assert c["ready"] and c["has_journal"] and c["issues"] == [] and c["unmapped"] == []
    for col in ("unadjusted", "adjusted"):
        assert c["tb_totals"][col]["diff"] == 0
    assert c["bs_totals"]["current"]["diff"] == 0 and c["bs_totals"]["prior"]["diff"] == 0
    assert c["bs_totals"]["reported_current"]["diff"] == 0
    assert all(a["unadjusted_vs_journal"] == 0 for a in c["accounts"].values())
    assert all(g["diff"] == 0 for g in c["journal"].values())
    fs = {l["name"]: l for l in c["fs_lines"]}
    cash = fs["현금및현금성자산"]
    assert cash["accounts"] == ["10301", "10302"] and cash["fs_vs_tb_current"] == 0 and cash["current"] == 20_500_000
    # 감사인 수정분개 A001: 급여 300,000 / 미지급비용 300,000
    assert c["accounts"]["80100"]["adjustment"] == 300_000 and c["accounts"]["80100"]["adjusted"] == 3_300_000
    assert c["accounts"]["26200"]["adjustment"] == 300_000
    assert fs["기타유동부채"]["reported_current"] == 0 and fs["기타유동부채"]["current"] == 300_000
    # 이익잉여금: 회사 21,000,000(당기순이익 7,000,000 포함) → 수정후 20,700,000
    re_ = fs["이익잉여금"]
    assert re_["includes_net_income"] and re_["fs_vs_tb_current"] == 0
    assert re_["adjustment"] == -300_000 and re_["current"] == 20_700_000
    assert c["income"]["net_income_reported"] == 7_000_000 and c["income"]["net_income"] == 6_700_000


def test_change_rate():
    assert engine.change_rate(5_500_000, 15_000_000) == 36.7
    assert engine.change_rate(-2_500_000, 15_000_000) == -16.7
    assert engine.change_rate(300_000, 0) is None  # 전기 0원: 증감률 없음(화면은 '신규')
    assert engine.change_rate(None, 100) is None and engine.change_rate(100, None) is None
    assert engine.change_rate(100, -200) == 50.0
    c = engine.compute_ledger(_ledger())
    assert c["accounts"]["10301"]["change_rate"] == -16.7 and c["accounts"]["10302"]["change_rate"] is None
    cash = next(l for l in c["fs_lines"] if l["name"] == "현금및현금성자산")
    assert cash["change_rate"] == 36.7 and cash["tb_sums"]["change"] == 5_500_000
    assert c["bs_totals"]["change"]["자산"] == 8_500_000 and c["bs_totals"]["change_rate"]["자산"] == 28.3
    assert c["income"]["net_income_change"] == 1_700_000 and c["income"]["net_income_change_rate"] == 34.0
    assert export._chg(-2_500_000, -16.7, 15_000_000) == "증감 ▼ 2,500,000 (▼ 16.7%)"
    assert export._chg(8_000_000, None, 0) == "증감 ▲ 8,000,000 (신규)"


def test_blank_amount_is_not_zero():
    L = _ledger()
    _acct(L, "10302")["unadjusted"] = ""
    c = engine.compute_ledger(L)
    assert c["accounts"]["10302"]["adjusted"] is None and c["accounts"]["10302"]["change"] is None
    assert c["tb_totals"]["adjusted"]["debit"] is None
    cash = next(l for l in c["fs_lines"] if l["name"] == "현금및현금성자산")
    assert cash["current"] is None and cash["fs_vs_tb_current"] is None
    assert any("미입력 금액 1칸" in m for m in _msgs(c["issues"], "warn"))


def test_unmapped_and_mismatch_issues():
    L = _ledger()
    _acct(L, "10302")["fs_line"] = "단기금융상품"  # 재무제표에 없는 항목
    c = engine.compute_ledger(L)
    assert c["unmapped"] == ["10302"] and c["accounts"]["10302"]["adjusted"] is None
    assert c["tb_totals"]["unadjusted"]["debit"] is None  # 분류 못 한 계정이 있으면 차대 합계를 내지 않음
    m = _msgs(c["issues"])
    assert any("10302 정기예금: 표시 항목 '단기금융상품'이 재무제표에 없음" in x for x in m)
    assert any("재무제표 '현금및현금성자산' 당기 금액과 시산표 합계가 8,000,000원 다름" in x for x in m)
    L = _ledger()
    L["adjustments"].append({"id": "a2", "je_no": "A002", "date": "", "description": "", "lines": [{"code": "10301", "debit": 100, "credit": ""}]})
    m = _msgs(engine.compute_ledger(L)["issues"], "block")
    assert any("전표 A002: 차변 100원과 대변 0원이 다름" in x for x in m)
    L["adjustments"][-1]["lines"].append({"code": "99999", "debit": "", "credit": 100})
    assert any("시산표에 없는 계정코드 '99999'" in x for x in _msgs(engine.compute_ledger(L)["issues"]))


def test_without_journal_and_without_files():
    L = _ledger()
    L["journal"] = []
    c = engine.compute_ledger(L)
    assert not c["has_journal"] and c["accounts"]["10301"]["unadjusted_vs_journal"] is None and c["issues"] == []
    g = engine.account_ledger(L, "80100")
    assert g["base"] == 3_000_000 and [ln["kind"] for ln in g["lines"]] == ["수정분개"] and g["closing_vs_adjusted"] == 0
    e = engine.compute_ledger({})
    assert not e["ready"] and _msgs(e["issues"]) == ["시산표를 올리지 않음", "재무제표를 올리지 않음"]


def test_account_ledger_with_journal():
    g = engine.account_ledger(_ledger(), "10301")
    assert g["base"] == 15_000_000
    assert [ln["balance"] for ln in g["lines"]] == [42_000_000, 23_500_000, 20_500_000, 12_500_000]
    assert g["closing_vs_adjusted"] == 0
    g = engine.account_ledger(_ledger(), "25100")
    assert g["calc"]["side"] == -1 and g["closing"] == 7_500_000
    g = engine.account_ledger(_ledger(), "80100")
    assert [ln["kind"] for ln in g["lines"]] == ["일반", "수정분개"] and g["closing"] == 3_300_000


# ---------------------------------------------------------------- 양식 업로드
def test_template_round_trip():
    L = _ledger()
    fs = imports.parse_fs(imports.write_fs(L["fs"]))
    tb = imports.parse_tb(imports.write_tb(L["accounts"]))
    assert fs["errors"] == [] and tb["errors"] == []
    strip = lambda rows: [{k: v for k, v in r.items() if k != "loc"} for r in rows]  # noqa: E731
    assert strip(fs["rows"]) == strip(L["fs"]) and strip(tb["accounts"]) == strip(L["accounts"])
    assert fs["rows"][0]["loc"] == "재무제표!2행"


def _xlsx(sheet, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_errors():
    bad = _xlsx("재무제표", [imports.FS_COLUMNS, ["재무상태표", "수익", "A", 1, 2], ["손익계산서", "비용", "", 1, 2],
                            ["재무상태표", "자산", "B", "1.5", None], ["재무상태표", "자산", "B", 1, 1]])
    r = imports.parse_fs(bad)
    assert len(r["errors"]) == 4 and r["rows"][0]["current"] == ""  # 빈 금액은 0이 아니라 빈칸
    r = imports.parse_tb(_xlsx("시산표", [imports.TB_COLUMNS, ["101", "현금", "", None, 10], ["101", "현금2", "", 1, 1]]))
    assert r["accounts"][0]["prior"] == "" and any("두 번" in e for e in r["errors"])
    try:
        imports.parse_tb(_xlsx("시산표", [["코드", "이름"]]))
        assert False
    except imports.ImportError_ as e:
        assert "열 이름" in str(e)


def test_upload_flow_api():
    cl = TestClient(app)
    cid = cl.post("/api/companies", json={"name": "업로드(가상)"}).json()["id"]
    eid = cl.post("/api/engagements", json={"company_id": cid, "period_end": "2025-12-31"}).json()["id"]
    r = cl.get(f"/api/engagements/{eid}/ledger").json()
    assert not r["computed"]["ready"] and r["procedures"]
    assert cl.get("/api/templates/fs.xlsx").status_code == 200
    L = _ledger()
    bad = cl.post(f"/api/engagements/{eid}/upload/fs", content=b"not xlsx")
    assert bad.status_code == 400
    r = cl.post(f"/api/engagements/{eid}/upload/fs", content=imports.write_fs(L["fs"]), headers={"X-Filename": "fs.xlsx"}).json()
    assert len(r["ledger"]["fs"]) == 9 and not r["computed"]["ready"]
    accts = copy.deepcopy(L["accounts"])
    accts[1]["fs_line"] = ""
    r = cl.post(f"/api/engagements/{eid}/upload/tb", content=imports.write_tb(accts)).json()
    assert r["computed"]["ready"] and r["computed"]["unmapped"] == ["10302"]
    r = cl.post(f"/api/engagements/{eid}/mapping", json={"10302": "현금및현금성자산"}).json()
    assert r["computed"]["unmapped"] == [] and r["computed"]["issues"] == []
    assert not r["computed"]["has_journal"]
    r = cl.post(f"/api/engagements/{eid}/adjustments", json={"adjustments": L["adjustments"]}).json()
    assert r["computed"]["accounts"]["80100"]["adjusted"] == 3_300_000
    assert next(c for c in cl.get("/api/tree").json() if c["id"] == cid)["engagements"][0]["has_ledger"]
    # 계정과목을 골라 조서 만들기
    wid = cl.post("/api/workpapers", json={"engagement_id": eid, "procedure_id": "cash.schedule_recon",
                                          "linked_accounts": ["10301", "10302"]}).json()["id"]
    got = cl.get(f"/api/workpapers/{wid}").json()
    assert got["data"]["linked_accounts"] == ["10301", "10302"]
    assert got["computed"]["linked"]["unadjusted"] == 20_500_000
    assert cl.get(f"/api/engagements/{eid}/ledger").json()["engagement"]["workpapers"][0]["linked_accounts"] == ["10301", "10302"]


# ---------------------------------------------------------------- 조서 참조
def test_workpaper_link_checks(resolved):
    d = records.default_data(resolved, "2025-12-31")
    assert d["linked_accounts"] == []
    c = engine.compute(resolved, d, None)
    assert any("정산표가 없음" in i["message"] for i in c["open_items"] if i["level"] == "info")
    c = engine.compute(resolved, d, _ledger())
    assert any("계정과목을 선택하지 않음" in i["message"] for i in c["open_items"] if i["level"] == "warn")
    d["linked_accounts"] = ["10301", "10302"]
    d["recon"]["gl_balance"] = "20500000"
    c = engine.compute(resolved, d, _ledger())
    assert c["linked"]["unadjusted"] == 20_500_000 and c["linked"]["change_rate"] == 36.7
    assert not any("정산표" in i["category"] for i in c["open_items"] if i["level"] == "warn")
    d["linked_accounts"] = ["10301"]
    c = engine.compute(resolved, d, _ledger())
    assert any("회사 시산표 당기 합계 12,500,000원과 입력한 총계정원장 잔액 20,500,000원이 다름" in i["message"] for i in c["open_items"])
    assert d["recon"]["gl_balance"] == "20500000"  # 참조값은 입력값을 바꾸지 않는다


# ---------------------------------------------------------------- 가상 사례·엑셀
def test_sample_consistency():
    L = samples.ledger()
    assert _acct(L, "10301")["unadjusted"] == 12_500_000 and _acct(L, "10302")["unadjusted"] == 8_000_000
    j006 = next(ln for ln in L["journal"] if ln["je_no"] == "J006" and ln["code"] == "10302")
    k02 = next(r for r in samples.classification(content.resolve(samples.CLS_BINDING))["classify"]["rows"] if r["item_id"] == "K02")
    assert j006["date"] == k02["acquired_date"]


def test_seeded(seeded):
    cl = TestClient(app)
    has = {c["name"]: c["engagements"][0]["has_ledger"] for c in cl.get("/api/tree").json()}
    assert has == {"가상 차이(주)": True, "가상 정상(주)": False, "가상 자료부족(주)": False}
    eid = records.load_workpaper(seeded[1])["engagement_id"]
    r = cl.get(f"/api/engagements/{eid}/ledger").json()
    assert r["computed"]["ready"] and r["engagement"]["company_id"]
    assert {w["wp_no"]: w["linked_accounts"] for w in r["engagement"]["workpapers"]} == {"A-200": ["10301", "10302"], "A-210": ["10301", "10302"]}
    assert cl.get(f"/api/engagements/{eid}/ledger/accounts/10302").json()["lines"][0]["je_no"] == "J006"
    assert cl.get(f"/api/engagements/{eid}/ledger/accounts/00000").status_code == 404
    other = records.load_workpaper(seeded[0])["engagement_id"]
    assert not cl.get(f"/api/engagements/{other}/ledger").json()["computed"]["ready"]
    assert cl.get(f"/api/engagements/{other}/ledger/accounts/10301").status_code == 404
    assert cl.get(f"/api/companies/{r['engagement']['company_id']}").json()["name"] == "가상 차이(주)"


def test_export_shows_linked_accounts(seeded):
    def text(wid):
        data, _ = export.build(records.load_workpaper(wid))
        return "\n".join(str(c.value) for row in load_workbook(io.BytesIO(data))["조서"].iter_rows() for c in row if c.value is not None)
    t = text(seeded[1])
    assert "10302 정기예금 (현금및현금성자산)" in t and "참조값" in t and "증감 ▲ 8,000,000 (신규)" in t
    assert "이 감사건에 정산표 없음" in text(seeded[0])
