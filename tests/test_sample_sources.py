"""가상 회사 자료 엑셀(samples/sources)이 가상 사례 조서의 근거 위치·금액과 맞는지 확인한다."""
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import samples

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import make_sample_sources as mss  # noqa: E402

FOLDER = {company: case for case, company in mss.CASES.items()}


@pytest.fixture(scope="module")
def sources(tmp_path_factory):
    out = tmp_path_factory.mktemp("sources")
    mss.build(out)
    return out


def _cell(folder, loc):
    """'명세_v1 명세!E4' → 해당 파일·시트·셀 값."""
    file, ref = loc.split(" ", 1)
    sheet, addr = ref.split("!")
    return load_workbook(folder / f"{file}.xlsx")[sheet][addr].value


def _locs(data):
    """조서 기록에서 (근거 위치, 금액) 쌍을 모은다. 금액이 빈칸이면 위치도 없다."""
    pairs = []
    for st in data["recon"]["sets"]:
        pairs.append((st["stated_total_loc"], st["stated_total"]))
        for r in st["rows"]:
            pairs.append((r["schedule_loc"], r["schedule_amount"]))
            pairs.append((r["subledger_loc"], r["subledger_amount"]))
    pairs.append((data["recon"]["gl_loc"], data["recon"]["gl_balance"]))
    return pairs


@pytest.mark.parametrize("company,builder", samples.CASES, ids=[c for c, _ in samples.CASES])
def test_locations_match_sample_amounts(sources, resolved, company, builder):
    folder = sources / FOLDER[company]
    data = builder(resolved)
    checked = 0
    for loc, amount in _locs(data):
        if amount == "":
            assert loc == "", f"금액이 빈칸인데 위치가 있음: {loc}"
            continue
        assert _cell(folder, loc) == int(amount), loc
        checked += 1
    assert checked >= 4


@pytest.mark.parametrize("company,builder", samples.CASES, ids=[c for c, _ in samples.CASES])
def test_files_follow_receipt_status(sources, resolved, company, builder):
    """입수한 자료는 파일이 있고, 요청만 한 자료는 파일이 없다."""
    folder = sources / FOLDER[company]
    for m in builder(resolved)["materials"]:
        assert (folder / m["location"]).exists() == (m["receipt"] == "입수"), m["location"]


def test_every_file_marked_virtual(sources):
    for path in sources.rglob("*.xlsx"):
        ws = load_workbook(path).active
        texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert mss.NOTICE in texts, path.name
