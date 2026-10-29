"""공통 계정 분류(절차 연결 기준)와, 콘텐츠에서 읽는 엑셀 서식 위치·근거 문구."""
import copy
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import content, engine, export, imports, records, samples
from app.main import app

G = content.load_account_groups()


def _row(name, statement="재무상태표", category="자산", **kw):
    return {"name": name, "statement": statement, "category": category, **kw}


def test_auto_match_ignores_spacing_and_needs_same_statement():
    assert engine.match_group(_row("현금 및 현금성자산"), G) == ("cash", "자동")
    assert engine.match_group(_row("현금및현금등가물"), G) == ("cash", "자동")
    assert engine.match_group(_row("이익잉여금(결손금)", category="자본"), G) == ("retained_earnings", "자동")
    assert engine.match_group(_row("현금및현금성자산", "손익계산서", "수익"), G) == (None, "없음")
    assert engine.match_group(_row("이름 모를 항목"), G) == (None, "없음")


def test_manual_override():
    assert engine.match_group(_row("예치금", group="short_term_financial"), G) == ("short_term_financial", "직접")
    assert engine.match_group(_row("현금및현금성자산", group=""), G) == (None, "직접")  # 직접 '분류 없음'


def test_net_income_line_by_group():
    L = copy.deepcopy(samples.ledger())
    del L["net_income_line"]
    next(r for r in L["fs"] if r["name"] == "이익잉여금")["name"] = "이익잉여금(결손금)"
    for a in L["accounts"]:
        if a["fs_line"] == "이익잉여금":
            a["fs_line"] = "이익잉여금(결손금)"
    c = engine.compute_ledger(L, G)
    assert c["net_income_line"] == "이익잉여금(결손금)" and c["issues"] == []
    assert engine.compute_ledger(L)["net_income_line"] is None  # 분류 목록 없이는 이름으로 찍지 않는다
    c = engine.compute_ledger(samples.ledger(), G)
    assert c["accounts"]["10302"]["group"] == "cash"


def test_content_validation(tmp_content):
    p = tmp_content / "bindings" / "cash.schedule_recon.json"
    b = json.loads(p.read_text(encoding="utf-8"))
    b["account_groups"] = ["cash", "없는분류"]
    b["form_refs"] = ["FORM4000:A!A99"]
    b["basis"]["refs"].append("KSA2026:999.1")
    p.write_text(json.dumps(b, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(content.ContentError) as e:
        content.resolve("cash.schedule_recon")
    errs = "\n".join(e.value.errors)
    assert "공통 계정 분류에 없는 id '없는분류'" in errs
    assert "FORM4000:A!A99" in errs and "KSA2026:999.1" in errs


def _text(wp):
    data, _ = export.build(wp)
    return "\n".join(str(c.value) for row in load_workbook(io.BytesIO(data))["조서"].iter_rows() for c in row if c.value is not None)


def test_export_text_comes_from_content(seeded):
    t = _text(records.load_workpaper(seeded[0]))
    assert "감사조서서식_4000 계정별 실증절차 (KIFRS용) 2026 — 서식 A!A17 (2. 현금및현금성자산 및 장단기금융상품 명세서)" in t
    assert "[출처: 서식 A!J17]" in t
    assert "감사기준서 330.20 · 감사기준서 330.30 · 감사기준서 500.9 (회계감사기준 전문(全文) 2026년 7월 개정)" in t
    t = _text(records.load_workpaper(seeded[3]))
    assert "미확인 — 감사조서서식_4000에서 이 절차의 위치를 확인하지 못함" in t and "[출처 미확인 — 사용자가 정함]" in t
    assert "분류 기준(기업회계기준서 제1007호·제1001호) 문단은 미확인" in t


def test_export_old_content_without_new_keys(seeded):
    wp = records.load_workpaper(seeded[0])
    for k in ("form_refs", "form_note", "basis"):
        wp["content"].pop(k, None)
    t = _text(wp)
    assert "서식 A!A17 (2. 현금및현금성자산 및 장단기금융상품 명세서)" in t and "이전 콘텐츠 버전" in t


def test_groups_api_and_procedure_matching():
    cl = TestClient(app)
    cid = cl.post("/api/companies", json={"name": "분류(가상)"}).json()["id"]
    eid = cl.post("/api/engagements", json={"company_id": cid, "period_end": "2025-12-31"}).json()["id"]
    L = samples.ledger()
    fs = copy.deepcopy(L["fs"])
    fs[0]["name"] = "현금 및 현금등가물"
    accts = copy.deepcopy(L["accounts"])
    for a in accts[:2]:
        a["fs_line"] = "현금 및 현금등가물"
    cl.post(f"/api/engagements/{eid}/upload/fs", content=imports.write_fs(fs))
    r = cl.post(f"/api/engagements/{eid}/upload/tb", content=imports.write_tb(accts)).json()
    line = r["computed"]["fs_lines"][0]
    assert (line["group"], line["group_source"]) == ("cash", "자동")
    assert {p["id"] for p in r["procedures"] if "cash" in p["account_groups"]} == {"cash.schedule_recon", "cash.classification_review"}
    r = cl.post(f"/api/engagements/{eid}/groups", json={"현금 및 현금등가물": ""}).json()
    assert r["computed"]["fs_lines"][0]["group"] is None and r["computed"]["accounts"]["10301"]["group"] is None
    r = cl.post(f"/api/engagements/{eid}/groups", json={"현금 및 현금등가물": None}).json()
    assert r["computed"]["fs_lines"][0]["group_source"] == "자동"
    assert cl.post(f"/api/engagements/{eid}/groups", json={"현금 및 현금등가물": "nope"}).status_code == 400
    wid = cl.post("/api/workpapers", json={"engagement_id": eid, "procedure_id": "cash.schedule_recon", "linked_accounts": ["10301"]}).json()["id"]
    assert cl.get(f"/api/workpapers/{wid}").json()["account_group_of"]["10301"] == "cash"
