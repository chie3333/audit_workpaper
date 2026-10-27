"""웹(엔진) 계산과 엑셀 수식 계산이 같은지, 출력 내용이 기록과 맞는지 확인한다.
LibreOffice가 있어야 수식 값을 확인할 수 있다."""
import copy
import io
import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import engine, export, records

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from recalc_xlsx import recalc  # noqa: E402

needs_lo = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice 없음")

KEYS = {"recalc": "recalc_total", "stated_vs_recalc": "stated_vs_recalc", "sub_total": "subledger_total",
        "gl": "gl_balance", "subledger_vs_gl": "subledger_vs_gl", "stated": "stated_total"}


def _excel_values(tmp_path, wp, name):
    data, cellmap = export.build(wp)
    src = tmp_path / f"{name}.xlsx"
    src.write_bytes(data)
    out = recalc(src, tmp_path / "out")
    return load_workbook(out, data_only=True), cellmap


def _norm(v):
    return None if v == "미입력" else v


def _assert_same(wb, cellmap, computed):
    for sid, m in cellmap["sets"].items():
        c = computed["sets"][sid]
        main = wb["조서"]
        for k, cell in m["summary"].items():
            assert _norm(main[cell].value) == c[KEYS[k]], (sid, k, main[cell].value, c[KEYS[k]])
        rows_ws = wb[m["sheet_rows"]]
        diffs = [_norm(rows_ws.cell(r, 7).value) for r in range(m["rows"]["first"], m["rows"]["last"] + 1)]
        assert diffs == [c["rows"][rid]["diff"] for rid in c["rows"]], (sid, diffs)


@needs_lo
def test_web_equals_excel_for_samples(seeded, tmp_path):
    for wid in seeded:
        wp = records.load_workpaper(wid)
        wb, cm = _excel_values(tmp_path, wp, f"wp{wid}")
        _assert_same(wb, cm, engine.compute(wp["content"], wp["data"]))


@needs_lo
def test_detail_sheet_for_long_table(seeded, tmp_path):
    wid = seeded[1]
    wp = records.load_workpaper(wid)
    d = copy.deepcopy(wp["data"])
    rows = d["recon"]["sets"][0]["rows"]
    for i in range(40):
        rows.append({"id": f"x{i}", "item_id": f"Z{i:02d}", "name": "가상", "schedule_amount": str(1000 * i),
                     "subledger_amount": "" if i == 7 else str(1000 * i - (5 if i == 3 else 0)), "schedule_loc": "", "subledger_loc": ""})
    records.save_workpaper(wid, d, base_rev=wp["rev"])
    wp = records.load_workpaper(wid)
    wb, cm = _excel_values(tmp_path, wp, "long")
    assert "검토표_상세" in wb.sheetnames
    assert cm["sets"]["s_orig"]["sheet_rows"] == "검토표_상세"
    _assert_same(wb, cm, engine.compute(wp["content"], wp["data"]))


def _text(wb):
    ws = wb["조서"]
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_output_marks_status_and_sample(seeded):
    wp = records.load_workpaper(seeded[1])  # 차이 사례
    data, _ = export.build(wp)
    t = _text(load_workbook(io.BytesIO(data)))
    assert "가상 데이터" in t
    assert "검토 상태: 미검토" in t
    assert "미처리 발견 사항 2건" in t
    assert "증거 파일은 이 엑셀에 포함되어 있지 않다" in t
    assert "작성하지 않음 — 연결 검토 상태: 확인 중" in t  # 수행하지 않은 명시 문장을 만들지 않음
    assert "다운로드는 검토 완료를 뜻하지 않는다" in t or "엑셀 다운로드는 검토 완료를 뜻하지 않는다" in t


def test_output_original_and_revised(seeded):
    wp = records.load_workpaper(seeded[2])  # 자료 부족 사례
    data, _ = export.build(wp)
    t = _text(load_workbook(io.BytesIO(data)))
    assert "현재 대사: 수정본 명세서" in t
    assert "이전 대사 기록(수정 전, 보존용): 원본 명세서" in t
    assert "수정본 (원본: 명세_v1)" in t
    assert "미입수 자료: 보조원장_K02" in t


def test_output_uses_selected_revision(seeded):
    wid = seeded[0]
    wp = records.load_workpaper(wid)
    d = copy.deepcopy(wp["data"])
    d["fields"]["conclusion"] = "새 결론"
    records.save_workpaper(wid, d, base_rev=wp["rev"])
    old = export.build(records.load_workpaper(wid, wp["rev"]))[0]
    t = _text(load_workbook(io.BytesIO(old)))
    assert "새 결론" not in t and "최신 저장본 아님" in t
