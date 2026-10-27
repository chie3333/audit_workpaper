"""분류 검토 절차(cash.classification_review): 콘텐츠, 기간·합계 계산, 남은 작업, 저장, 엑셀, 가상 사례."""
import copy
import io
import shutil
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import content, engine, export, records, samples
from app.main import app

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import make_sample_sources as mss  # noqa: E402
from recalc_xlsx import recalc  # noqa: E402

BID = "cash.classification_review"
needs_lo = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice 없음")


@pytest.fixture
def cres():
    return content.resolve(BID)


def _row(rid="c1", **kw):
    r = {"id": rid, "item_id": "K02", "name": "가상", "product_type": "정기예금", "amount": "8000000", "amount_loc": "명세!E5",
         "maturity_type": "만기 있음", "acquired_date": "2025-11-15", "maturity_date": "2026-05-15", "terms_loc": "증서!C5:C6",
         "restricted": "없음", "restriction_note": "", "restriction_loc": "회신", "company_class": "현금및현금성자산",
         "auditor_class": "단기금융상품", "basis": "증서 확인", "note": ""}
    r.update(kw)
    return r


def _data(cres, rows):
    d = records.default_data(cres, "2025-12-31")
    d["classify"]["rows"] = rows
    return d


def _msgs(c):
    return [i["message"] for i in c["open_items"] if i["level"] == "block"]


# ---------------------------------------------------------------- 콘텐츠
def test_content_resolves_with_cl_and_unverified_basis(cres):
    assert cres["assertions"]["default"] == ["CL"]
    assert cres["assertions"]["unverified"] is True and cres["assertions"]["refs"] == []
    assert cres["required_documents"] == ["schedule", "terms", "restriction"]
    assert cres["classify"]["classes"] == ["현금및현금성자산", "단기금융상품", "장기금융상품"]
    # 분류 기준(기업회계기준서) 안내는 출처 없이 미확인으로만 표시한다
    why = next(s for s in cres["steps"] if s["id"] == "classify")["guidance"][0]
    assert why["type"] == "보완 설명" and why["unverified"] and not why["refs"]
    assert all(g["review"]["status"] == "AI 초안" for s in cres["steps"] for g in s["guidance"])


def test_content_requires_classes(tmp_content):
    import json
    p = tmp_content / "bindings" / f"{BID}.json"
    b = json.loads(p.read_text(encoding="utf-8"))
    b["classify"]["classes"] = []
    p.write_text(json.dumps(b, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(content.ContentError) as e:
        content.resolve(BID)
    assert any("분류 선택지" in m for m in e.value.errors)


def test_default_data_has_classify_not_recon(cres):
    d = records.default_data(cres, "2025-12-31")
    assert d["classify"] == {"schedule_material_id": "", "rows": []}
    assert "recon" not in d and "scan" not in d and "list_check" not in d["fields"]


# ---------------------------------------------------------------- 계산
def test_days_reclass_and_totals(cres):
    d = _data(cres, [_row("c1", item_id="K01", product_type="보통예금", amount="12500000", maturity_type="만기 없음",
                          acquired_date="", maturity_date="", auditor_class="현금및현금성자산"),
                     _row("c2")])
    c = engine.compute(cres, d)["classify"]
    assert c["rows"]["c1"]["acq_days"] is None and c["rows"]["c1"]["no_maturity"]
    assert c["rows"]["c1"]["reclass"] == 0
    assert c["rows"]["c2"]["acq_days"] == 181
    assert c["rows"]["c2"]["remain_days"] == 135
    assert c["rows"]["c2"]["reclass"] == 8_000_000
    t = c["totals"]
    assert t["현금및현금성자산"] == {"company": 20_500_000, "auditor": 12_500_000, "diff": -8_000_000}
    assert t["단기금융상품"] == {"company": 0, "auditor": 8_000_000, "diff": 8_000_000}
    assert t["장기금융상품"]["diff"] == 0


def test_blank_amount_or_class_stops_totals(cres):
    c = engine.compute(cres, _data(cres, [_row("c1"), _row("c2", amount="")]))["classify"]
    assert c["total_amount"] is None and c["totals"]["단기금융상품"]["company"] is None
    assert c["rows"]["c2"]["reclass"] is None
    c = engine.compute(cres, _data(cres, [_row("c1", auditor_class="")]))["classify"]
    assert c["totals"]["단기금융상품"]["company"] == 0  # 회사 분류는 다 있음
    assert c["totals"]["단기금융상품"]["auditor"] is None and c["totals"]["단기금융상품"]["diff"] is None


def test_program_does_not_classify(cres):
    """기간이 짧아도 감사인 판단 분류를 채우거나 차이를 만들지 않는다."""
    d = _data(cres, [_row("c1", acquired_date="2025-12-01", maturity_date="2026-01-31", auditor_class="", basis="")])
    c = engine.compute(cres, d)
    assert c["classify"]["rows"]["c1"]["auditor"] == "" and c["diffs"] == []
    assert any("감사인 판단 분류 미입력" in m for m in _msgs(c))


# ---------------------------------------------------------------- 남은 작업
def test_checks(cres):
    d = _data(cres, [_row("c1", restricted="미확인", basis=""),
                     _row("c2", item_id="K03", restricted="있음", restriction_note=""),
                     _row("c3", item_id="K04", maturity_type="만기 있음", maturity_date="", acquired_date="2025-01-01"),
                     _row("c4", item_id="K05", acquired_date="2026-01-01", maturity_date="2025-06-30")])
    m = _msgs(engine.compute(cres, d))
    assert any("K02: 사용제한 여부 미확인" in x for x in m)
    assert any("K02: 감사인 판단 분류의 근거 미기록" in x for x in m)
    assert any("K03: 사용제한 내용 미기록" in x for x in m)
    assert any("K04: 만기일 미입력" in x for x in m)
    assert any("K05: 만기일이 취득일보다 빠름" in x for x in m)
    assert any("자료 미등록: 예금증서·상품 약정서" in x for x in m)
    assert any("자료 미등록: 사용제한 확인 자료" in x for x in m)
    assert not any("총계정원장" in x or "보조부" in x for x in m)


def test_reclass_needs_finding_and_link_state(cres):
    d = _data(cres, [_row("c1")])
    c = engine.compute(cres, d)
    assert [x["key"] for x in c["diffs"]] == ["cls:row:c1"]
    assert any("발견 사항 기록 없음" in x for x in _msgs(c))
    d["findings"] = [{"id": "f1", "link": "cls:row:c1", "title": "t", "status": "미해결"}]
    c = engine.compute(cres, d)
    assert not any("발견 사항 기록 없음" in x for x in _msgs(c))
    assert c["ui"]["findings"]["f1"]["link_state"] == "open"
    d["classify"]["rows"][0]["auditor_class"] = "현금및현금성자산"
    assert engine.compute(cres, d)["ui"]["findings"]["f1"]["link_state"] == "zero_or_blank"
    d["classify"]["rows"] = []
    assert engine.compute(cres, d)["ui"]["findings"]["f1"]["link_state"] == "missing"


# ---------------------------------------------------------------- 저장·API
def test_api_create_preview_save():
    cl = TestClient(app)
    assert BID in [p["id"] for p in cl.get("/api/procedures").json()]
    cid = cl.post("/api/companies", json={"name": "분류(가상)"}).json()["id"]
    eid = cl.post("/api/engagements", json={"company_id": cid, "period_end": "2025-12-31"}).json()["id"]
    wid = cl.post("/api/workpapers", json={"engagement_id": eid, "procedure_id": BID}).json()["id"]
    wp = cl.get(f"/api/workpapers/{wid}").json()
    d = wp["data"]
    d["classify"]["rows"] = [_row("c1")]
    c = cl.post(f"/api/workpapers/{wid}/preview", json={"data": d}).json()
    assert c["classify"]["rows"]["c1"]["acq_days"] == 181
    saved = cl.post(f"/api/workpapers/{wid}/save", json={"data": d, "base_rev": wp["rev"]}).json()
    assert saved["data"]["classify"]["rows"][0]["maturity_date"] == "2026-05-15"
    assert cl.get(f"/api/workpapers/{wid}/export.xlsx").status_code == 200


# ---------------------------------------------------------------- 가상 사례
def test_sample_case(seeded):
    wid = seeded[3]
    wp = records.load_workpaper(wid)
    assert wp["company"] == "가상 차이(주)" and wp["content"]["id"] == BID
    c = engine.compute(wp["content"], wp["data"])
    assert c["classify"]["rows"]["c2"]["acq_days"] == 181
    assert [d["key"] for d in c["diffs"]] == ["cls:row:c2"] and c["diffs"][0]["has_finding"]
    assert c["status"]["findings_unresolved"] == 1


def test_sample_source_matches(tmp_path, cres):
    mss.build(tmp_path)
    row = next(r for r in samples.classification(cres)["classify"]["rows"] if r["item_id"] == "K02")
    file, ref = row["terms_loc"].split(" ", 1)
    sheet, rng = ref.split("!")
    ws = load_workbook(tmp_path / "차이" / f"{file}.xlsx")[sheet]
    start, end = rng.split(":")
    assert ws[start].value.date().isoformat() == row["acquired_date"]
    assert ws[end].value.date().isoformat() == row["maturity_date"]
    file, ref = row["amount_loc"].split(" ", 1)
    sheet, addr = ref.split("!")
    assert load_workbook(tmp_path / "차이" / f"{file}.xlsx")[sheet][addr].value == int(row["amount"])


# ---------------------------------------------------------------- 엑셀
def _text(wb):
    return "\n".join(str(c.value) for row in wb["조서"].iter_rows() for c in row if c.value is not None)


def test_export_content(seeded):
    wp = records.load_workpaper(seeded[3])
    data, cm = export.build(wp)
    wb = load_workbook(io.BytesIO(data))
    t = _text(wb)
    assert "5. 분류 검토표" in t and "가상 데이터" in t
    assert "미확인 — 감사조서서식_4000에서 이 절차의 위치를 확인하지 못함" in t
    assert "분류(CL)" in t and "출처 미확인" in t
    assert "명세서 스캔" not in t and "재무제표 정보와 회계기록의 연결" not in t
    ws = wb["조서"]
    assert str(ws[cm["classify"]["rows"]["c2"]["acq_days"]].value).startswith("=IF(")
    assert set(cm["classify"]["totals"]) == set(wp["content"]["classify"]["classes"])


def _n(v):
    return None if v in ("미입력", "해당 없음") else v


@needs_lo
def test_web_equals_excel(seeded, tmp_path):
    wid = seeded[3]
    wp = records.load_workpaper(wid)
    d = copy.deepcopy(wp["data"])
    d["classify"]["rows"].append(_row("c9", item_id="K09", amount="", auditor_class=""))
    for case in (wp["data"], d):
        w = dict(wp, data=case)
        data, cm = export.build(w)
        src = tmp_path / f"cls{len(case['classify']['rows'])}.xlsx"
        src.write_bytes(data)
        ws = load_workbook(recalc(src, tmp_path / "out"), data_only=True)["조서"]
        c = engine.compute(wp["content"], case)["classify"]
        for rid, cells in cm["classify"]["rows"].items():
            for k, cell in cells.items():
                assert _n(ws[cell].value) == c["rows"][rid][k], (rid, k)
        for k, cells in cm["classify"]["totals"].items():
            for kk, cell in cells.items():
                assert _n(ws[cell].value) == c["totals"][k][kk], (k, kk)
