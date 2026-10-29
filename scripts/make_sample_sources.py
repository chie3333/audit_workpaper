"""가상 사례 3종이 근거 위치로 가리키는 회사 자료(명세서·보조원장·총계정원장)를 엑셀로 만든다.

전부 가상 데이터다. 회사·금융기관·계좌·금액은 실재하지 않는다.
셀 위치는 app/samples.py의 근거 위치(예: '명세_v1 명세!E4')와 같아야 한다 — tests/test_sample_sources.py가 대조한다.
명세서의 표시 합계는 회사가 적어 준 값이라 수식이 아니라 값으로 넣는다(차이 사례는 일부러 틀린 합계).

사용: python scripts/make_sample_sources.py [출력폴더]   (기본: samples/sources)
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app import imports, samples  # noqa: E402
DEFAULT_OUT = ROOT / "samples" / "sources"

PERIOD = "2025-12-31"
NOTICE = "가상 데이터 — 회사·금융기관·계좌·금액은 실재하지 않음. 감사조서 프로그램 연습용."
NUM = "#,##0"
BOLD = Font(bold=True)
THIN = Side(style="thin")
LINE = Border(top=THIN, bottom=THIN)

K01 = ("K01", "가나은행(가상)", "보통예금")
K02 = ("K02", "다라은행(가상)", "정기예금")

# 사례 폴더 → 회사명 (app/samples.py의 CASES와 같은 회사)
CASES = {
    "정상": "가상 정상(주)",
    "차이": "가상 차이(주)",
    "자료부족": "가상 자료부족(주)",
}


def _head(ws, title, company, cols, widths):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"회사: {company}    기준일: {PERIOD}    단위: 원"
    for i, (name, width) in enumerate(zip(cols, widths)):
        c = ws.cell(row=3, column=i + 1, value=name)
        c.font = BOLD
        c.border = LINE
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = width


def _notice(ws, row):
    ws.cell(row=row, column=1, value=NOTICE).font = Font(italic=True, color="888888")


def schedule(path, company, k01, k02, stated_total, version):
    """명세서: 명세!E4·E5 계좌별 금액, E6 표시 합계."""
    wb = Workbook()
    ws = wb.active
    ws.title = "명세"
    _head(ws, f"현금및현금성자산·장단기금융상품 명세서 ({version})", company,
          ["번호", "계좌", "금융기관", "종류", "금액"], [6, 8, 18, 12, 16])
    for i, ((item, bank, kind), amount) in enumerate([(K01, k01), (K02, k02)]):
        r = 4 + i
        for col, v in enumerate([i + 1, item, bank, kind, amount], start=1):
            ws.cell(row=r, column=col, value=v)
        ws.cell(row=r, column=5).number_format = NUM
    ws["A6"] = "합계"
    ws["A6"].font = BOLD
    ws["E6"] = stated_total
    ws["E6"].number_format = NUM
    ws["E6"].font = BOLD
    for col in range(1, 6):
        ws.cell(row=6, column=col).border = LINE
    _notice(ws, 8)
    wb.save(path)


def subledger(path, company, rows):
    """예금별 보조원장 잔액: 잔액!D4부터 계좌별 기말잔액."""
    wb = Workbook()
    ws = wb.active
    ws.title = "잔액"
    _head(ws, "예금별 보조원장 기말잔액", company, ["계좌", "금융기관", "종류", "기말잔액"], [8, 18, 12, 16])
    for i, ((item, bank, kind), amount) in enumerate(rows):
        r = 4 + i
        for col, v in enumerate([item, bank, kind, amount], start=1):
            ws.cell(row=r, column=col, value=v)
        ws.cell(row=r, column=4).number_format = NUM
    _notice(ws, 4 + len(rows) + 2)
    wb.save(path)


def general_ledger(path, company, balance):
    """총계정원장: 원장!D4 현금및현금성자산 기말잔액."""
    wb = Workbook()
    ws = wb.active
    ws.title = "원장"
    _head(ws, "총계정원장 계정별 잔액", company, ["계정코드", "계정과목", "차대", "기말잔액"], [10, 20, 6, 16])
    for col, v in enumerate(["V101", "현금및현금성자산", "차", balance], start=1):
        ws.cell(row=4, column=col, value=v)
    ws["D4"].number_format = NUM
    _notice(ws, 6)
    wb.save(path)


def deposit_certificate(path, company, item, bank, product, acquired, maturity, amount, early_exit, pledge):
    """예금증서: 증서!C5 가입일(취득일), C6 만기일, C7 금액."""
    wb = Workbook()
    ws = wb.active
    ws.title = "증서"
    ws["A1"] = f"{product} 예금증서 사본"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"예금주: {company}"
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 40
    for r, (label, value) in enumerate([("계좌", item), ("금융기관", bank), ("가입일", acquired), ("만기일", maturity),
                                        ("금액(원)", amount), ("중도해지", early_exit), ("질권 설정", pledge)], start=3):
        ws.cell(row=r, column=2, value=label).font = BOLD
        c = ws.cell(row=r, column=3, value=value)
        c.alignment = Alignment(horizontal="left")
        if label in ("가입일", "만기일"):
            c.number_format = "yyyy-mm-dd"
        if label == "금액(원)":
            c.number_format = NUM
    _notice(ws, 11)
    wb.save(path)


def build(out: Path = DEFAULT_OUT) -> list[Path]:
    made = []

    def folder(case):
        d = out / case
        d.mkdir(parents=True, exist_ok=True)
        return d

    # 정상: 명세서·보조원장·총계정원장 모두 일치
    d, co = folder("정상"), CASES["정상"]
    schedule(d / "명세_v1.xlsx", co, 12_500_000, 8_000_000, 20_500_000, "v1")
    subledger(d / "보조원장.xlsx", co, [(K01, 12_500_000), (K02, 8_000_000)])
    general_ledger(d / "총계정원장.xlsx", co, 20_500_000)

    # 차이: 명세서 K01을 12,400,000으로 잘못 적고 표시 합계는 20,500,000
    d, co = folder("차이"), CASES["차이"]
    schedule(d / "명세_v1.xlsx", co, 12_400_000, 8_000_000, 20_500_000, "v1")
    subledger(d / "보조원장.xlsx", co, [(K01, 12_500_000), (K02, 8_000_000)])
    general_ledger(d / "총계정원장.xlsx", co, 20_500_000)
    # 차이 회사의 분류 검토 조서(A-210)가 쓰는 K02 예금증서
    deposit_certificate(d / "예금증서_K02.xlsx", co, *K02, date(2025, 11, 15), date(2026, 5, 15), 8_000_000,
                        "가능 (약정이율보다 낮은 중도해지이율 적용)", "없음")

    # 차이 회사의 감사건 자료 입력용: 회사 재무제표·시산표(감사 전), 앱 업로드 양식 그대로
    led = samples.ledger()
    (d / "재무제표.xlsx").write_bytes(imports.write_fs(led["fs"], notice=NOTICE))
    (d / "시산표.xlsx").write_bytes(imports.write_tb(led["accounts"], notice=NOTICE))

    # 자료 부족: 원본·수정본 명세서, K01 보조원장만 입수. K02 보조원장은 미입수라 만들지 않는다.
    d, co = folder("자료부족"), CASES["자료부족"]
    schedule(d / "명세_v1.xlsx", co, 12_400_000, 8_000_000, 20_500_000, "v1")
    schedule(d / "명세_v2.xlsx", co, 12_500_000, 8_000_000, 20_500_000, "v2 수정")
    subledger(d / "보조원장_K01.xlsx", co, [(K01, 12_500_000)])
    general_ledger(d / "총계정원장.xlsx", co, 20_500_000)

    for case in CASES:
        made += sorted((out / case).glob("*.xlsx"))
    return made


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    for p in build(out):
        print(p.relative_to(out.parent) if out.parent in p.parents else p)
